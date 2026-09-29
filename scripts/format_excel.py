import sys
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension

src, dst = sys.argv[1], sys.argv[2]
wb = load_workbook(src)

thin = Side(style='thin')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
header_fill = PatternFill('solid', fgColor='D9EAF7')
header_font = Font(bold=True)

for ws in wb.worksheets:
    ws.freeze_panes = 'A2'
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = ws.dimensions if ws.max_row > 1 and ws.max_column > 1 else None
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = border
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 24
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 0
        for cell in ws[letter]:
            value = '' if cell.value is None else str(cell.value)
            max_len = max(max_len, max((len(x) for x in value.split('\n')), default=0))
        # Keep sheets readable without creating extremely wide columns.
        width = min(max(max_len + 2, 10), 45)
        if ws.title == 'Dokumen Lengkap' and col == 2:
            width = min(max(max_len + 2, 25), 90)
        ws.column_dimensions[letter].width = width
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 30 if ws.max_column <= 2 else 22

# Make the main data sheet easier to scan.
if 'Data' in wb.sheetnames:
    ws = wb['Data']
    if ws.max_column == 1:
        ws.column_dimensions['A'].width = 90
    elif ws.max_column <= 8:
        for c in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(c)].width = min(max(ws.column_dimensions[get_column_letter(c)].width or 10, 15), 40)

wb.save(dst)
