import sys
from docx import Document

out, title = sys.argv[1], sys.argv[2]
text = open(sys.argv[3], 'r', encoding='utf-8', errors='replace').read()
doc = Document()
doc.add_heading(title or 'Dokumen', level=1)
for line in text.replace('\r\n','\n').replace('\r','\n').split('\n'):
    doc.add_paragraph(line)
doc.save(out)
print(out)
